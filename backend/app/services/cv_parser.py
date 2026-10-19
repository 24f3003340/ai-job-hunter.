"""Rule-based CV text -> structured draft.

Design points:
* Parsing is deterministic string processing. Document text is never executed or sent to a model
  here, so embedded instructions cannot steer it. Lines that look like instructions to an AI are
  removed from the parse and reported.
* Nothing is trusted. Every uncertain field becomes a flag; the draft only reaches the profile
  through the user's confirm step.
* An education entry's status is never promoted to "completed" on a guess: it is only set to
  `in_progress` (explicit marker or future end date); otherwise it is left null with a suggestion.
"""
import io
import re
import unicodedata
from datetime import date

import docx
from pypdf import PdfReader

PARSER_VERSION = "1.0"

_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
_DATE = (r"(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{4}"
         r"|\d{1,2}/\d{4}|\d{4}-\d{2}|\d{4})")
_NOW = r"(?:present|current|currently|ongoing|now|till date|to date)"
_RANGE = re.compile(rf"(?P<s>{_DATE})\s*(?:-|–|—|to)\s*(?:expected\s+)?(?P<e>{_DATE}|{_NOW})", re.I)
_EXPECTED = re.compile(rf"(?:expected|exp\.?|pursuing)\s*(?:in|:)?\s*(?P<e>{_DATE})", re.I)

_SECTIONS = {
    "summary": {"summary", "profile", "objective", "about me", "career objective", "professional summary"},
    "education": {"education", "academic background", "academics", "qualifications", "educational qualifications",
                  "educational qualification", "academic qualifications"},
    "skills": {"skills", "technical skills", "key skills", "core competencies", "skills and tools", "technologies",
               "skills & tools", "tech stack"},
    "experience": {"experience", "work experience", "professional experience", "employment", "employment history",
                   "internships", "internship", "internship experience", "work history"},
    "projects": {"projects", "academic projects", "personal projects", "key projects"},
    "certifications": {"certifications", "certificates", "licenses", "certifications & courses", "courses",
                       "certifications and courses"},
    "other": {"achievements", "awards", "interests", "languages", "hobbies", "references", "declaration",
              "extracurricular", "extracurricular activities", "positions of responsibility"},
}
_HEADING = {alias: sec for sec, aliases in _SECTIONS.items() for alias in aliases}

_DEGREE = re.compile(r"\b(b\.?\s?tech|m\.?\s?tech|b\.?\s?e\.?|m\.?\s?e\.?|b\.?\s?sc|m\.?\s?sc|b\.?\s?c\.?a|m\.?\s?c\.?a|"
                     r"mba|bba|b\.?\s?com|m\.?\s?com|ph\.?\s?d|bachelor|master|diploma|b\.?\s?a\.?|m\.?\s?a\.?|"
                     r"high school|senior secondary|secondary|12th|10th|hsc|ssc|class\s+(?:xii|xi|x|12|10))\b", re.I)
_INSTITUTION = re.compile(r"\b(university|institute|college|school|academy|iit|nit|iiit|polytechnic)\b", re.I)
_IN_PROGRESS = re.compile(r"\b(expected|pursuing|persuing|currently|ongoing|in progress|present|final year|"
                          r"appearing|studying|3rd year|2nd year|4th year|third year|second year)\b", re.I)
_COMPLETED = re.compile(r"\b(graduated|completed|awarded|passed out)\b", re.I)
_TITLE_WORDS = re.compile(r"\b(engineer|developer|intern|analyst|scientist|manager|lead|associate|trainee|consultant|"
                          r"researcher|designer|architect|administrator|executive|assistant|programmer|tester|"
                          r"specialist|officer|head|director)\b", re.I)
_BULLET = re.compile(r"^\s*(?:[•●▪■◦\-*–·]|\d+[.)])\s+")
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_PHONE = re.compile(r"(?<!\d)(\+?\d[\d\s().-]{8,}\d)(?!\d)")

_INJECTION = re.compile(
    r"ignore (?:all |any |the )?(?:previous|prior|above|earlier) (?:instructions|prompts?)|"
    r"disregard (?:all |any |the )?(?:previous|prior|above|earlier)|"
    r"system prompt|you are (?:now )?(?:an? )?(?:ai|assistant|language model|chatgpt|claude)|"
    r"as an ai\b|<\|?(?:system|im_start|assistant)|\[/?inst\]|"
    r"(?:rate|rank|score) (?:this|the) (?:candidate|resume|cv)|hire this candidate|"
    r"do not (?:flag|mention|report)|new instructions?:", re.I)

_ZERO_WIDTH = dict.fromkeys(map(ord, "​‌‍⁠﻿­"), None)


class Unreadable(Exception):
    """Raised when no usable text can be extracted (scanned/image-only or corrupt document)."""


# ---- text extraction --------------------------------------------------------------------------

def extract_text(kind: str, data: bytes) -> str:
    try:
        if kind == "pdf":
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        else:
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(parts)
    except Exception as exc:  # parser libraries raise many types on hostile input
        raise Unreadable(f"Could not read the document ({type(exc).__name__})") from exc
    text = normalize(text)
    if len(re.sub(r"\s", "", text)) < 50:
        raise Unreadable("Very little text could be extracted. The file may be a scanned image; OCR is not "
                         "supported, so please enter your details manually or upload a text-based CV.")
    return text


# PDF/Word bullets often extract as control or private-use glyphs (e.g. Symbol font U+F0B7).
_LINE_BULLET = re.compile("(?m)^[ \\t]*[\x7f\x95\uf0b7\uf0a7\uf076\uf0d8\uf0fc\uf06e\uf02d\u25cf\u25aa\u25a0]+[ \\t]*")


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).translate(_ZERO_WIDTH)
    text = _LINE_BULLET.sub("• ", text)
    text = "".join(ch for ch in text if ch in "\n\t" or unicodedata.category(ch)[0] != "C")
    return re.sub(r"[ \t]+", " ", text)


# ---- helpers ----------------------------------------------------------------------------------

def _norm_date(tok: str) -> str | None:
    tok = tok.strip().lower().rstrip(".")
    if re.fullmatch(_NOW, tok):
        return None
    if m := re.fullmatch(r"(\d{4})-(\d{2})", tok):
        return f"{m[1]}-{m[2]}" if 1 <= int(m[2]) <= 12 else m[1]
    if m := re.fullmatch(r"(\d{1,2})/(\d{4})", tok):
        return f"{m[2]}-{int(m[1]):02d}" if 1 <= int(m[1]) <= 12 else m[2]
    if m := re.fullmatch(r"([a-z]+)\.?\s+(\d{4})", tok):
        mon = _MONTHS.get(m[1][:3])
        return f"{m[2]}-{mon:02d}" if mon else m[2]
    if re.fullmatch(r"\d{4}", tok):
        return tok
    return None


def _is_now(tok: str) -> bool:
    return bool(re.fullmatch(_NOW, tok.strip(), re.I))


def _split_range(text: str):
    """Return (start, end, current, text_without_range)."""
    if m := _RANGE.search(text):
        cur = _is_now(m["e"])
        return _norm_date(m["s"]), (None if cur else _norm_date(m["e"])), cur, (text[:m.start()] + text[m.end():])
    if m := _EXPECTED.search(text):
        return None, _norm_date(m["e"]), False, text
    return None, None, False, text


def _clean(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip(" \t-–—|,;:()")


def _is_future(end: str | None, today: date) -> bool:
    if not end:
        return False
    y, _, m = end.partition("-")
    return (int(y), int(m or 12)) > (today.year, today.month)


def _is_past(end: str | None, today: date) -> bool:
    """True only when we are sure the end date has passed (a bare current year is ambiguous)."""
    if not end:
        return False
    y, _, m = end.partition("-")
    if not m:
        return int(y) < today.year
    return (int(y), int(m)) < (today.year, today.month)


class _Flags(list):
    def add(self, code: str, path: str, severity: str, message: str) -> None:
        self.append({"id": f"{code}:{path}", "code": code, "path": path, "severity": severity, "message": message})


def _split_sections(lines: list[str]):
    sections: dict[str, list[str]] = {"header": []}
    cur = "header"
    for line in lines:
        key = re.sub(r"[:\s]+$", "", line.strip().lower())
        key = re.sub(r"\s+", " ", key)
        if key in _HEADING and len(line.strip()) < 45:
            cur = _HEADING[key]
            sections.setdefault(cur, [])
            continue
        sections[cur].append(line)
    return sections


# ---- section parsers --------------------------------------------------------------------------

def _parse_education(lines, flags: _Flags, today: date):
    entries: list[dict] = []
    cur: dict | None = None

    def finish():
        nonlocal cur
        if cur is None:
            return
        i = len(entries)
        text = cur.pop("_text")
        start, end, current, _ = _split_range(text)
        cur["start_date"], cur["end_date"] = start, end
        exp = _EXPECTED.search(text)
        if exp:
            cur["end_date"] = _norm_date(exp["e"])
        in_prog = bool(_IN_PROGRESS.search(text)) or current or _is_future(cur["end_date"], today)
        done = bool(_COMPLETED.search(text))
        cur["status"], cur["suggested_status"], cur["confidence"] = None, None, "medium"
        if in_prog and not done:
            cur["status"], cur["confidence"] = "in_progress", "high"
        elif done and not in_prog:
            cur["suggested_status"] = "completed"
            flags.add("education_status_unconfirmed", f"education[{i}]", "warning",
                      "Completion is stated in the document; please confirm the status explicitly.")
        elif in_prog and done:
            flags.add("education_status_conflict", f"education[{i}]", "warning",
                      "The entry contains both in-progress and completed wording; choose the correct status.")
        elif _is_past(cur["end_date"], today):
            cur["suggested_status"] = "completed"
            flags.add("education_status_inferred", f"education[{i}]", "warning",
                      "End date is in the past but completion is not stated; confirm whether this degree is completed.")
        else:
            flags.add("education_status_unknown", f"education[{i}]", "warning",
                      "Could not determine whether this education is completed or in progress.")
        if not cur.get("institution"):
            cur["confidence"] = "low"
            flags.add("education_institution_missing", f"education[{i}]", "warning", "Institution not found.")
        entries.append(cur)
        cur = None

    for raw in lines:
        line = _BULLET.sub("", raw).strip()
        if not line:
            continue
        if _DEGREE.search(line) and (cur is None or cur.get("degree")):
            finish()
            cur = {"degree": "", "institution": "", "_text": ""}
        if cur is None:
            cur = {"degree": "", "institution": "", "_text": ""}
        cur["_text"] += " " + line
        _, _, _, rest = _split_range(line)
        rest = _clean(_EXPECTED.sub("", rest))
        if _DEGREE.search(line) and not cur["degree"]:
            parts = [p for p in re.split(r"\s+[|–—]\s+|,\s+(?=[A-Z])", rest) if p.strip()]
            deg = [p for p in parts if _DEGREE.search(p)]
            cur["degree"] = _clean(deg[0] if deg else rest)
            inst = [p for p in parts if _INSTITUTION.search(p)]
            if inst:
                cur["institution"] = _clean(inst[0])
        elif _INSTITUTION.search(line) and not cur["institution"]:
            cur["institution"] = _clean(rest)
    finish()
    return entries


def _entries(lines):
    """Group lines into [header_lines, bullet_lines] blocks."""
    blocks: list[dict] = []
    cur: dict | None = None
    for raw in lines:
        if not raw.strip():
            continue
        is_bullet = bool(_BULLET.match(raw))
        text = _BULLET.sub("", raw).strip()
        if is_bullet:
            if cur is None:
                cur = {"header": [], "bullets": []}
                blocks.append(cur)
            cur["bullets"].append(text)
        elif cur is not None and cur["bullets"] and (text[:1].islower() or len(text) > 90):
            cur["bullets"][-1] += " " + text  # wrapped bullet line
        elif cur is None or cur["bullets"]:
            cur = {"header": [text], "bullets": []}
            blocks.append(cur)
        else:
            cur["header"].append(text)
    return blocks


def _parse_experience(lines, flags: _Flags):
    out = []
    for i, b in enumerate(_entries(lines)):
        header = " ; ".join(b["header"])
        start, end, current, rest = _split_range(header)
        parts = [p for p in (_clean(x) for x in re.split(r"\s*;\s*|\s+\|\s+|\s+@\s+|\s+at\s+|\s+[–—-]\s+", rest)) if p]
        title = employer = ""
        conf = "medium"
        if len(parts) >= 2:
            tp = [p for p in parts if _TITLE_WORDS.search(p)]
            if tp:
                title = tp[0]
                employer = next((p for p in parts if p != title), "")
            else:
                title, employer, conf = parts[0], parts[1], "low"
        elif len(parts) == 1:
            if _TITLE_WORDS.search(parts[0]):
                title = parts[0]
            else:
                employer = parts[0]
            conf = "low"
        else:
            conf = "low"
        if not title or not employer:
            flags.add("experience_header_unclear", f"experience[{i}]", "warning",
                      "Could not reliably identify both the job title and the employer.")
        if not (start or end or current):
            flags.add("experience_dates_missing", f"experience[{i}]", "warning", "No dates found for this entry.")
        out.append({"title": title, "employer": employer, "start_date": start, "end_date": end, "current": current,
                    "description": "", "bullets": b["bullets"], "confidence": conf})
    return out


def _parse_projects(lines, flags: _Flags):
    out = []
    for i, b in enumerate(_entries(lines)):
        header = b["header"]
        name_line = header[0] if header else ""
        _, _, _, name_line = _split_range(name_line)
        tech: list[str] = []
        m = re.search(r"\s+[|–—-]\s+(.+)$|\(([^)]+)\)\s*$", name_line)
        if m:
            tech = [t.strip() for t in re.split(r"[,/;]", m[1] or m[2]) if t.strip()]
            name_line = name_line[:m.start()]
        desc = []
        for h in header[1:]:
            tm = re.match(r"(?:tech(?:nologies| stack)?|tools|built with)\s*:\s*(.+)", h, re.I)
            if tm:
                tech += [t.strip() for t in re.split(r"[,;]", tm[1]) if t.strip()]
            else:
                desc.append(h)
        name = _clean(name_line)
        if not name:
            flags.add("project_name_missing", f"projects[{i}]", "warning", "Project name not found.")
        out.append({"name": name, "description": " ".join(desc), "bullets": b["bullets"], "technologies": tech,
                    "url": None, "confidence": "medium" if name else "low"})
    return out


def _parse_skills(lines, flags: _Flags):
    names, seen = [], set()
    for raw in lines:
        line = _BULLET.sub("", raw).strip()
        if ":" in line and len(line.split(":", 1)[0]) <= 30:
            line = line.split(":", 1)[1]
        for tok in re.split(r"[,;|•·]", line):
            tok = _clean(tok)
            if not tok:
                continue
            if len(tok) > 40 or len(tok.split()) > 5:
                flags.add("skill_unclear", f"skills[{len(names)}]", "info", f"Skipped long skills text: {tok[:50]}...")
                continue
            if tok.lower() not in seen:
                seen.add(tok.lower())
                names.append({"name": tok, "level": None, "verified": False})
    return names


def _parse_certs(lines, flags: _Flags):
    out = []
    for raw in lines:
        line = _BULLET.sub("", raw).strip()
        if not line:
            continue
        m = re.search(_DATE, line, re.I)
        dt = _norm_date(m[0]) if m else None
        rest = _clean(re.sub(_DATE, "", line, flags=re.I)) if m else line
        parts = re.split(r"\s+[|–—-]\s+|\s+by\s+", rest, maxsplit=1)
        out.append({"name": _clean(parts[0]), "issuer": _clean(parts[1]) if len(parts) > 1 else "", "date": dt})
    return out


# ---- entry point ------------------------------------------------------------------------------

def parse_text(text: str, today: date | None = None) -> tuple[dict, list[dict]]:
    today = today or date.today()
    flags = _Flags()
    kept = []
    for n, line in enumerate(text.split("\n")):
        if _INJECTION.search(line):
            flags.add("suspicious_instruction", f"line[{n}]", "warning",
                      f"Text resembling an instruction to an AI system was ignored: “{line.strip()[:80]}”")
        else:
            kept.append(line)
    secs = _split_sections(kept)
    header = [ln for ln in secs["header"] if ln.strip()]

    name = ""
    for ln in header[:5]:
        s = ln.strip()
        if s and not re.search(r"[@\d/]|http|www\.", s) and len(s.split()) <= 5:
            name = s
            flags.add("name_from_first_line", "full_name", "info", "Name was taken from the first line; please check it.")
            break
    if not name:
        flags.add("name_missing", "full_name", "warning", "Could not identify your name.")
    blob = " ".join(header[:12] + [ln for ln in kept[:15]])
    email = (_EMAIL.search(blob) or [None])[0]
    phone = (_PHONE.search(blob) or [None])[0]

    parsed = {
        "full_name": name, "email": email, "phone": _clean(phone) if phone else "",
        "education": _parse_education(secs.get("education", []), flags, today),
        "experience": _parse_experience(secs.get("experience", []), flags),
        "projects": _parse_projects(secs.get("projects", []), flags),
        "skills": _parse_skills(secs.get("skills", []), flags),
        "certifications": _parse_certs(secs.get("certifications", []), flags),
    }
    for sec in ("education", "skills", "experience", "projects"):
        if not parsed[sec]:
            flags.add("section_missing", sec, "info", f"No {sec} section was detected.")
    return parsed, list(flags)
