import io
from xml.sax.saxutils import escape

import docx
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer


def _lines(content: dict, claims: list[dict], state: dict) -> list[tuple[str, str]]:
    """Flatten to (kind, text). Rejected claims are omitted; unresolved ones are marked."""
    status = {c["id"]: c["status"] for c in claims}
    decisions = state["decisions"]

    def claim_text(c) -> str | None:
        if status.get(c["id"]) == "needs_review":
            d = decisions.get(c["id"])
            if d == "reject":
                return None
            if d != "accept":
                return "[UNVERIFIED] " + c["text"]
        return c["text"]

    def dates(e) -> str:
        s, en = e.get("start_date"), "Present" if e.get("current") else e.get("end_date")
        return f"{s or ''} – {en or ''}".strip(" –") if (s or en) else ""

    out: list[tuple[str, str]] = []
    h = content["header"]
    out.append(("name", h.get("name") or ""))
    out.append(("contact", " | ".join(x for x in (h.get("email"), h.get("phone")) if x)))
    if content.get("summary") and (t := claim_text(content["summary"])):
        out += [("heading", "Summary"), ("para", t)]
    if content["skills"]:
        out += [("heading", "Skills"), ("para", ", ".join(s["name"] for s in content["skills"]))]
    if content["experience"]:
        out.append(("heading", "Experience"))
        for e in content["experience"]:
            out.append(("sub", " — ".join(x for x in (e.get("title"), e.get("employer"), dates(e)) if x)))
            out += [("bullet", t) for b in e["bullets"] if (t := claim_text(b))]
    if content["projects"]:
        out.append(("heading", "Projects"))
        for p in content["projects"]:
            tech = f" ({', '.join(p['technologies'])})" if p.get("technologies") else ""
            out.append(("sub", (p.get("name") or "") + tech))
            out += [("bullet", t) for b in p["bullets"] if (t := claim_text(b))]
    if content["education"]:
        out.append(("heading", "Education"))
        for e in content["education"]:
            st = "Expected / in progress" if e["status"] == "in_progress" else "Completed"
            out.append(("para", " — ".join(x for x in (e.get("degree"), e.get("institution"), dates(e), st) if x)))
    if content["certifications"]:
        out.append(("heading", "Certifications"))
        out += [("bullet", " — ".join(x for x in (c["name"], c.get("issuer"), c.get("date")) if x))
                for c in content["certifications"]]
    return out


def to_docx(content: dict, claims: list[dict], state: dict) -> bytes:
    d = docx.Document()
    for kind, text in _lines(content, claims, state):
        if kind == "name":
            d.add_heading(text, 0)
        elif kind == "heading":
            d.add_heading(text, 2)
        elif kind == "bullet":
            d.add_paragraph(text, style="List Bullet")
        elif kind == "sub":
            d.add_paragraph().add_run(text).bold = True
        else:
            p = d.add_paragraph(text)
            p.paragraph_format.space_after = Pt(4)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_pdf(content: dict, claims: list[dict], state: dict) -> bytes:
    ss = getSampleStyleSheet()
    styles = {"name": ss["Title"], "contact": ss["Normal"], "heading": ss["Heading2"], "para": ss["Normal"],
              "sub": ParagraphStyle("sub", parent=ss["Normal"], fontName="Helvetica-Bold", spaceBefore=4)}
    flow, bullets = [], []

    def flush():
        nonlocal bullets
        if bullets:
            flow.append(ListFlowable([ListItem(Paragraph(escape(t), ss["Normal"])) for t in bullets],
                                     bulletType="bullet", leftIndent=14))
            bullets = []

    for kind, text in _lines(content, claims, state):
        if kind == "bullet":
            bullets.append(text)
            continue
        flush()
        flow.append(Paragraph(escape(text), styles[kind]))
    flush()
    flow.append(Spacer(1, 6))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.8 * cm,
                      bottomMargin=1.8 * cm, title="Resume").build(flow)
    return buf.getvalue()
